#!/usr/bin/env python3
"""Monta e valida a pauta editorial (Metodologia Carratu).

Modo 1, classificar (entrada CSV keyword,volume,kd,intent,source,cluster):
  python3 pauta_build.py post_data/keywords.csv --profile perfil.json --out post_data/pauta.json
  → cauda, intenção (heurística, se o CSV não trouxer), prioridade, cluster inicial, campos vazios para Claude.

Modo 2, validar e gerar (entrada pauta.json preenchido):
  python3 pauta_build.py post_data/pauta.json --profile perfil.json --validate --out-dir outputs
          [--start 2026-11-03] [--every 3] [--period 2026-11]
  → confere H1/title/H2/persona/duplicidade, distribui datas, escreve pauta-<blog>-<periodo>.md/.csv/.xlsx.
Copyright (c) 2026 Ubirajara Carratu, Carratu Digital. PolyForm Noncommercial 1.0.0.
"""
import argparse
import csv
import datetime as dt
import json
import re
import statistics
import sys
import unicodedata
from pathlib import Path

STOP = set("a o as os um uma uns umas de da do das dos em na no nas nos por para pra com sem e ou que se ao aos à às "
           "the of to and for in on with is are how what why".split())
INTENT_RULES = [
    ("transacional", r"\b(pre[çc]o|pre[çc]os|quanto custa|comprar|contratar|assinar|assinatura|plano|planos|teste gr[áa]tis|"
                     r"or[çc]amento|demo|demonstra[çc][ãa]o|cota[çc][ãa]o|pricing|buy|price|cost|trial)\b"),
    ("comercial", r"\b(melhor|melhores|vs|versus|comparativo|compara[çc][ãa]o|alternativa|alternativas|review|reviews|"
                  r"avalia[çc][ãa]o|vale a pena|top \d+|ranking|best|compare|comparison|alternatives)\b"),
    ("navegacional", r"\b(login|entrar|site oficial|app|aplicativo|baixar|download|telefone|contato|suporte)\b"),
    ("informacional", r"\b(o que [ée]|oque e|como|por que|porque|guia|exemplo|exemplos|significado|diferen[çc]a|passo a passo|"
                      r"dicas|tipos|quais|quando|onde|tutorial|what is|how to|guide|examples|tips)\b"),
]
GENERIC_H2 = {"introducao", "conclusao", "beneficios", "vantagens", "desvantagens", "dicas", "exemplos", "resumo"}


def norm(s):
    s = unicodedata.normalize("NFKD", s or "").encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]", " ", s)).strip()


def cw(s):
    return [w for w in norm(s).split() if w not in STOP and len(w) > 1]


MODIFIERS = set("como melhor melhores o que e preco precos quanto custa guia exemplo exemplos dicas tipos quais qual quando onde "
                "por que porque vs versus alternativa alternativas comparativo review vale pena obrigatorio login what how why best".split())


def tail(keyword, volume):
    n = len(cw(keyword))
    is_question = bool(re.match(r"^(o que|como|por que|porque|quanto|quando|qual|quais|onde|what|how|why|when)\b", norm(keyword)))
    if is_question or n >= 5:
        return "long"
    if volume is not None:
        if volume >= 1000:
            return "short" if n <= 2 else "medium"
        if volume < 100:
            return "long"
        return "medium"
    return "short" if n <= 1 else "medium" if n <= 3 else "long"


def intent(keyword, brand_terms=()):
    k = norm(keyword)
    if any(norm(b) and norm(b) in k for b in brand_terms):
        return "navegacional"
    for name, rx in INTENT_RULES:
        if re.search(rx, k):
            return name
    return "informacional"


def to_num(v):
    try:
        return float(str(v).replace(".", "").replace(",", ".")) if str(v).strip() not in ("", "None", "null") else None
    except ValueError:
        return None


def guess_cluster(keyword, freq):
    """Cluster = palavra de conteúdo (não modificador) mais frequente no universo de palavras-chave."""
    ws = [w for w in cw(keyword) if w not in MODIFIERS] or cw(keyword)
    return max(ws, key=lambda w: (freq.get(w, 0), -ws.index(w))) if ws else keyword


def classify(rows, profile):
    brand = [profile.get("marca", "")] + [c.split(".")[0] for c in profile.get("concorrentes", [])]
    vols = [to_num(r.get("volume")) for r in rows]
    vols = [v for v in vols if v is not None]
    med = statistics.median(vols) if vols else None
    freq = {}
    for r in rows:
        for w in set(cw(r["keyword"])) - MODIFIERS:
            freq[w] = freq.get(w, 0) + 1
    out = []
    for r in rows:
        kw = r["keyword"].strip()
        vol, kd = to_num(r.get("volume")), to_num(r.get("kd"))
        it = (r.get("intent") or "").strip().lower() or intent(kw, brand)
        tl = tail(kw, vol)
        cl = (r.get("cluster") or "").strip() or guess_cluster(kw, freq)
        score = 0
        if it == "comercial":
            score += 2
        elif it == "informacional":
            score += 1
        if med and vol and vol > med:
            score += 1
        if str(r.get("aio", "")).strip().lower() in ("1", "true", "sim", "yes"):
            score += 1
        if (kd is not None and kd < 30) or tl == "long":
            score += 1
        if it in ("transacional", "navegacional"):
            score -= 2
        if r.get("cannibal_risk"):
            score -= 2
        pr = 1 if score >= 4 else 2 if score >= 2 else 3
        out.append({
            "keyword": kw, "tail": tl, "intent": it, "volume": vol, "kd": kd,
            "aio": r.get("aio") or None, "source": r.get("source") or None, "cluster": cl,
            "priority": pr, "priority_score": score,
            "format": "pagina" if it == "transacional" else ("comparação" if it == "comercial" else ""),
            "funnel": {"informacional": "topo", "comercial": "meio", "transacional": "fundo", "navegacional": "fundo"}[it],
            "persona": "", "h1": "", "seo_title": "", "h2": [], "secondary_keywords": [],
            "cannibal_risk": r.get("cannibal_risk") or None, "publish_date": None, "notes": "",
        })
    out.sort(key=lambda x: (x["priority"], -(x["volume"] or 0)))
    return out


def validate(rows, profile, start, every):
    personas = {p["id"] for p in profile.get("personas", [])} | {p.get("nome", "") for p in profile.get("personas", [])}
    issues, seen = [], {}
    d = dt.date.fromisoformat(start) if start else None
    for i, r in enumerate(rows, 1):
        tag = f"L{i} \"{r['keyword']}\""
        h1, t = r.get("h1", ""), r.get("seo_title", "")
        if not h1:
            issues.append(f"{tag}: h1 vazio")
        else:
            if not 40 <= len(h1) <= 70:
                issues.append(f"{tag}: h1 com {len(h1)} caracteres (40 a 70)")
            kwords = cw(r["keyword"])
            first_third = norm(h1)[: max(20, len(norm(h1)) // 3 + 10)]
            if kwords and not all(any(w2.startswith(w[:max(4, len(w) - 2)]) for w2 in norm(h1).split()) for w in kwords):
                issues.append(f"{tag}: palavra-chave não aparece no h1")
            elif kwords and not any(w in first_third for w in kwords):
                issues.append(f"{tag}: palavra-chave fora do primeiro terço do h1 (aviso)")
        if not t:
            issues.append(f"{tag}: seo_title vazio")
        elif not 50 <= len(t) <= 60:
            issues.append(f"{tag}: seo_title com {len(t)} caracteres (50 a 60)")
        h2 = r.get("h2") or []
        if not 5 <= len(h2) <= 8:
            issues.append(f"{tag}: {len(h2)} H2 (5 a 8)")
        gen = [h for h in h2 if norm(h) in GENERIC_H2]
        if gen:
            issues.append(f"{tag}: H2 genéricos {gen}")
        if not any("pergunta" in norm(h) or "faq" in norm(h) for h in h2):
            issues.append(f"{tag}: falta H2 'Perguntas frequentes'")
        if r.get("persona") and personas and r["persona"] not in personas:
            issues.append(f"{tag}: persona \"{r['persona']}\" não está no perfil")
        if not r.get("persona"):
            issues.append(f"{tag}: persona vazia")
        sec = r.get("secondary_keywords") or []
        bad = [s for s in sec if not (isinstance(s, dict) and s.get("keyword") and s.get("intent"))]
        if bad:
            issues.append(f"{tag}: secondary_keywords precisa de objetos {{keyword, intent}}")
        key = (r["intent"], " ".join(sorted(cw(r["keyword"]))))
        if key in seen:
            issues.append(f"{tag}: mesma intenção e mesmas palavras que L{seen[key]} (fundir ou mudar ângulo)")
        seen.setdefault(key, i)
        if r.get("intent") == "transacional" and r.get("format") != "pagina":
            issues.append(f"{tag}: transacional costuma ser página, não post (aviso)")
        if d and not r.get("publish_date"):
            r["publish_date"] = d.isoformat()
            d += dt.timedelta(days=every)
    return issues


def write_outputs(rows, profile, out_dir, period):
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    base = out_dir / f"pauta-{profile.get('id', 'blog')}-{period}"
    cols = ["publish_date", "priority", "cluster", "keyword", "tail", "intent", "funnel", "format", "persona", "h1",
            "seo_title", "h2", "secondary_keywords", "volume", "kd", "aio", "cannibal_risk", "notes"]

    def flat(r, c):
        v = r.get(c)
        if c == "h2":
            return " | ".join(v or [])
        if c == "secondary_keywords":
            return " | ".join(f"{s['keyword']} ({s['intent']})" if isinstance(s, dict) else str(s) for s in (v or []))
        return "" if v is None else v

    with open(base.with_suffix(".csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(cols)
        for r in rows:
            w.writerow([flat(r, c) for c in cols])

    by = lambda k: {v: sum(1 for r in rows if r.get(k) == v) for v in sorted({r.get(k) for r in rows}, key=str)}
    md = [f"# Pauta {profile.get('marca', '')} · {period}", "",
          f"{len(rows)} artigos · intenção {by('intent')} · cauda {by('tail')} · personas {by('persona')}", "",
          "| Data | P | Cluster | Palavra-chave | Cauda | Intenção | Persona | H1 |", "|---|---|---|---|---|---|---|---|"]
    for r in rows:
        md.append(f"| {r.get('publish_date') or ''} | {r['priority']} | {r['cluster']} | {r['keyword']} | {r['tail']} | "
                  f"{r['intent']} | {r.get('persona', '')} | {r.get('h1', '')} |")
    md.append("")
    for i, r in enumerate(rows, 1):
        md += [f"## {i}. {r.get('h1') or r['keyword']}", "",
               f"- **Palavra-chave:** {r['keyword']} ({r['tail']}, {r['intent']}, {r['funnel']} de funil"
               + (f", volume {int(r['volume'])}" if r.get('volume') else "") + (f", KD {int(r['kd'])}" if r.get('kd') else "") + ")",
               f"- **Title:** {r.get('seo_title', '')}", f"- **Persona:** {r.get('persona', '')}",
               f"- **Formato:** {r.get('format', '')} · **Cluster:** {r['cluster']} · **Prioridade:** {r['priority']}",
               "- **Secundárias por intenção:** " + ", ".join(f"{s['keyword']} ({s['intent']})" if isinstance(s, dict) else str(s)
                                                              for s in r.get("secondary_keywords") or []),
               "- **Subtópicos (H2):**"] + [f"  {j}. {h}" for j, h in enumerate(r.get("h2") or [], 1)]
        if r.get("cannibal_risk"):
            md.append(f"- **Atenção, canibalização:** {r['cannibal_risk']}")
        if r.get("notes"):
            md.append(f"- Notas: {r['notes']}")
        md.append("")
    base.with_suffix(".md").write_text("\n".join(md), encoding="utf-8")

    xlsx = None
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "Pauta"
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="F43404")
        for r in rows:
            ws.append([flat(r, c) for c in cols])
        for col, width in zip("ABCDEFGHIJKLM", (12, 4, 16, 30, 8, 14, 8, 12, 14, 50, 45, 60, 50)):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"
        xlsx = base.with_suffix(".xlsx")
        wb.save(xlsx)
    except ImportError:
        pass
    return base, xlsx


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input")
    ap.add_argument("--profile", required=True)
    ap.add_argument("--out")
    ap.add_argument("--validate", action="store_true")
    ap.add_argument("--out-dir", default="outputs")
    ap.add_argument("--start")
    ap.add_argument("--every", type=int, default=3)
    ap.add_argument("--period", default=dt.date.today().strftime("%Y-%m"))
    a = ap.parse_args()
    profile = json.loads(Path(a.profile).read_text(encoding="utf-8"))

    if not a.validate:
        with open(a.input, encoding="utf-8") as f:
            rows = list(csv.DictReader(f))
        if not rows or "keyword" not in rows[0]:
            sys.exit("CSV precisa da coluna keyword (e, opcionalmente, volume,kd,intent,source,cluster,aio,cannibal_risk).")
        out = classify(rows, profile)
        data = {"blog": profile.get("id"), "rows": out}
        if a.out:
            Path(a.out).parent.mkdir(parents=True, exist_ok=True)
            Path(a.out).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        by = lambda k: {v: sum(1 for r in out if r[k] == v) for v in sorted({r[k] for r in out}, key=str)}
        print(f"{len(out)} palavras-chave · cauda {by('tail')} · intenção {by('intent')} · prioridade {by('priority')}")
        for r in out[:10]:
            print(f"  P{r['priority']} {r['tail']:6} {r['intent']:13} {r['keyword']}  [{r['cluster']}]")
        print("Preencha persona, h1, seo_title, h2[], secondary_keywords[{keyword,intent}] e rode com --validate.")
        return 0

    data = json.loads(Path(a.input).read_text(encoding="utf-8"))
    rows = data["rows"] if isinstance(data, dict) else data
    issues = validate(rows, profile, a.start, a.every)
    base, xlsx = write_outputs(rows, profile, a.out_dir, a.period)
    Path(a.input).write_text(json.dumps({"blog": profile.get("id"), "rows": rows}, ensure_ascii=False, indent=2), encoding="utf-8")
    hard = [i for i in issues if "(aviso)" not in i]
    print(f"{len(rows)} linhas · {len(hard)} problemas · {len(issues) - len(hard)} avisos")
    for i in issues:
        print("  -", i)
    print(f"Gerados: {base}.md, {base}.csv" + (f", {xlsx}" if xlsx else " (openpyxl ausente: sem .xlsx)"))
    return 1 if hard else 0


if __name__ == "__main__":
    sys.exit(main())
